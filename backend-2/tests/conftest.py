import io

import pytest
from docx import Document

SAMPLE_CV = [
    "Jan Novak",
    "jan@novak.sk | +421 905 123 456 | https://github.com/jan",
    "EXPERIENCE",
    "Backend Developer, Acme s.r.o. 03/2019 - present",
    "Built REST APIs in Node.js and Express, TypeScript, Docker, Azure SQL.",
    "Junior developer, Foo a.s. 2016 - 2018",
    "Python, Django, PostgreSQL. I react quickly and excel at teamwork.",
    "EDUCATION",
    "Ing. - Technical University of Kosice, Faculty of Electrical Engineering, 2011 - 2016",
    "Skills: JavaScript, React, Git, CI/CD, English",
]


def make_pdf(lines: list[str]) -> bytes:
    """Minimálne platné PDF s textom (Helvetica), bez externej knižnice."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = "BT /F1 10 Tf 40 800 Td 14 TL " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(lines: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def cv_pdf() -> bytes:
    return make_pdf(SAMPLE_CV)


@pytest.fixture
def cv_docx() -> bytes:
    return make_docx(SAMPLE_CV)
