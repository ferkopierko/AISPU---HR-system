"""Extrakcia surového textu z CV (PDF / DOCX)."""

import io
import zipfile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class UnsupportedFileError(ValueError):
    """Typ súboru nie je PDF ani DOCX."""


class FileParseError(ValueError):
    """Súbor je poškodený alebo z neho nejde prečítať text."""


def detect_type(filename: str, data: bytes) -> str:
    """Typ sa určí podľa obsahu (magic bytes), prípona musí sedieť."""
    name = (filename or "").lower()
    if data.startswith(b"%PDF-") and name.endswith(".pdf"):
        return "pdf"
    if data.startswith(b"PK\x03\x04") and name.endswith(".docx"):
        return "docx"
    raise UnsupportedFileError("Podporované sú iba súbory PDF a DOCX")


def extract_text(filename: str, data: bytes) -> str:
    kind = detect_type(filename, data)
    try:
        text = _pdf_text(data) if kind == "pdf" else _docx_text(data)
    except (PdfReadError, zipfile.BadZipFile, KeyError, ValueError) as err:
        raise FileParseError("Súbor sa nepodarilo prečítať (je poškodený?)") from err
    if not text.strip():
        raise FileParseError("V súbore sa nenašiel žiadny text (naskenované CV nie je podporované)")
    return text


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise FileParseError("PDF je zaheslované")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # zručnosti a prax sú v CV často v tabuľkách
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
